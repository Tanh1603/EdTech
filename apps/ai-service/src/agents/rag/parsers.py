import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


@dataclass(frozen=True)
class ParsedDocument:
    material_id: str
    title: str
    text: str


class TextParser:
    def parse(self, material_id: str, title: str, content: str) -> ParsedDocument:
        return ParsedDocument(material_id=material_id, title=title, text=content.strip())


class DocumentParser:
    def parse_bytes(
        self,
        material_id: str,
        title: str,
        content: bytes,
        mime_type: str | None = None,
        filename: str | None = None,
    ) -> ParsedDocument:
        mime = (mime_type or "").lower()
        name = (filename or "").lower()
        if "pdf" in mime or name.endswith(".pdf"):
            text = self._pdf_text(content)
        elif "wordprocessingml" in mime or name.endswith(".docx"):
            text = self._docx_text(content)
        elif "presentationml" in mime or name.endswith(".pptx"):
            text = self._pptx_text(content)
        else:
            text = content.decode("utf-8", errors="ignore")
        return ParsedDocument(material_id=material_id, title=title, text=_clean_text(text))

    def _pdf_text(self, content: bytes) -> str:
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    def _docx_text(self, content: bytes) -> str:
        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def _pptx_text(self, content: bytes) -> str:
        presentation = Presentation(BytesIO(content))
        lines: list[str] = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text:
                    lines.append(text)
        return "\n".join(lines)


def _clean_text(text: str) -> str:
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text)
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()
